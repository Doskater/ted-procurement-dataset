import csv
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ted_dataset.export import export_workbook, write_csv, COLUMNS, verify_workbook
from ted_dataset.qa import make_sample, review_sample
from ted_dataset.pipeline import process_snapshot, update_report
from ted_dataset.normalize import normalize
from ted_dataset.fetch import fetch_snapshot
from test_fetch import Session, page


def test_excel_roundtrip_types_and_formula_text(tmp_path, notice, refs):
    notice["title-proc"] = {"eng": '=HYPERLINK("https://evil.test","click")'}
    row = normalize(notice, refs, date(2026,8,6), date(2026,10,4)).row
    path = tmp_path / "test.xlsx"
    export_workbook(path, [row], [], {"selected": 1}, {"status": "pending"}, {"status": "not recorded"})
    wb = load_workbook(path)
    assert wb.sheetnames == ["Notices", "QA Summary", "Data Dictionary", "Rejected"]
    sheet = wb["Notices"]
    assert sheet.cell(2, COLUMNS.index("title") + 1).data_type == "s"
    assert sheet.cell(2, COLUMNS.index("cpv_code") + 1).data_type == "s"
    assert sheet.cell(2, COLUMNS.index("publication_date") + 1).is_date
    assert sheet.cell(2, COLUMNS.index("estimated_value") + 1).data_type == "n"
    assert sheet.cell(2, COLUMNS.index("source_url") + 1).hyperlink.target == row["source_url"]
    assert sheet.freeze_panes == "B2" and sheet.auto_filter.ref
    verify_workbook(path, [row])
    write_csv(tmp_path / "test.csv", COLUMNS, [row])
    with (tmp_path / "test.csv").open(encoding="utf-8-sig", newline="") as f:
        values = list(csv.DictReader(f))
    assert values[0]["title"] == row["title"]  # CSV has literal data; import as text.


def test_manual_pending_mismatch_and_complete(notice, refs):
    row = normalize(notice, refs, date(2026,8,6), date(2026,10,4)).row
    sample = make_sample([row], "snapshot", "dataset", {row["notice_id"]: "raw/page.json#/notices/0"})
    assert review_sample(sample, sample)["status"] == "pending"
    completed = [dict(sample[0], **{k: "match" for k in sample[0] if k.endswith("_match")},
                      review_status="reviewed", reviewed_by="Human", reviewed_at="2026-10-04")]
    assert review_sample(completed, sample)["status"] == "passed"
    completed[0]["title_match"] = "mismatch"
    assert review_sample(completed, sample)["status"] == "mismatch"
    completed[0]["review_status"] = "PASS"
    with pytest.raises(ValueError):
        review_sample(completed, sample)


@pytest.mark.parametrize("change", ["snapshot_id", "expected_title", "notice_id", "dataset_sha256"])
def test_manual_cannot_be_reused_for_other_data(notice, refs, change):
    row = normalize(notice, refs, date(2026,8,6), date(2026,10,4)).row
    expected = make_sample([row], "snapshot", "dataset", {})
    changed = [dict(expected[0], **{change: "tampered"})]
    with pytest.raises(ValueError):
        review_sample(changed, expected)


def test_manual_missing_or_duplicate_rows_rejected(notice, refs):
    row = normalize(notice, refs, date(2026,8,6), date(2026,10,4)).row
    expected = make_sample([row], "snapshot", "dataset", {})
    for supplied in ([], expected * 2):
        with pytest.raises(ValueError):
            review_sample(supplied, expected)


def test_offline_processing_repeatability_and_preservation(tmp_path, notice):
    run = tmp_path / "run"
    fetch_snapshot(run, date(2026,10,4), session=Session([page([notice]), page([])]))
    first, second = tmp_path / "first", tmp_path / "second"
    reference = Path(__file__).resolve().parents[1] / "data/reference"
    process_snapshot(run, first, reference)
    process_snapshot(run, second, reference)
    for name in ["validated_notices.csv", "valid_not_selected.csv", "record_ledger.csv", "rejected_or_uncertain.csv"]:
        assert (first / name).read_bytes() == (second / name).read_bytes()
    assert json.loads((first/"result.json").read_text())["metrics"] == json.loads((second/"result.json").read_text())["metrics"]
    manual = first / "qa/manual_qa_sample.csv"
    before = manual.read_bytes()
    update_report(run, first)
    assert manual.read_bytes() == before
    with pytest.raises(FileExistsError):
        process_snapshot(run, first, reference)
    assert "manual verification pending" in (first/"QA_REPORT.md").read_text()


def test_modified_dataset_blocks_report(tmp_path, notice):
    run, out = tmp_path / "run", tmp_path / "out"
    fetch_snapshot(run, date(2026,10,4), session=Session([page([notice]), page([])]))
    process_snapshot(run, out, Path(__file__).resolve().parents[1]/"data/reference")
    (out/"validated_notices.csv").write_text("tampered")
    with pytest.raises(ValueError, match="checksum"):
        update_report(run, out)
