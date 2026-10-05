"""Typed Excel delivery and exact, machine-readable CSV output."""
import csv
import math
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = ["notice_id", "title", "buyer", "buyer_country", "cpv_code", "cpv_description",
           "publication_date", "deadline", "estimated_value", "currency", "source_url",
           "qa_status", "qa_note", "notice_identifier", "notice_version", "notice_type",
           "title_language", "buyer_language", "lot_count"]
REJECT_COLUMNS = ["notice_id", "title", "source_url", "status", "raw_ref", "reason_code", "reason_detail"]
DICTIONARY = {
    "notice_id": "Unique TED publication-number. Text, not procedure ID.",
    "title": "Procedure title in source English if present, otherwise alphabetically first language code. No translation.",
    "buyer": "Sorted distinct names in chosen source language, joined with semicolons.",
    "buyer_country": "Sorted country codes from TED; not paired positionally with buyer names.",
    "cpv_code": "Single procedure-level main CPV, eight-digit text; prefixes 48 or 72.",
    "cpv_description": "Official English CPV 2008 description.",
    "publication_date": "Source calendar date. Timezone suffix validated without UTC day shifting.",
    "deadline": "Tender submission date for a single-lot notice only. Blank with QA note otherwise. No active-status guarantee.",
    "estimated_value": "Positive procedure-level estimate (estimated-value-proc). Blank if absent there, even when the sole lot has an estimate. No lot fallback, lot sums or currency conversion.",
    "currency": "Source ISO 4217 code paired with published value. Blank when value is absent.",
    "source_url": "Exact English detail URL returned by TED for this publication number.",
    "qa_status": "accepted means automatic record rules passed, not human review completed.",
    "qa_note": "Missing optional values and omitted ambiguous deadlines. MISSING_PROCEDURE_ESTIMATED_VALUE means absent at procedure level; it does not assert that lot estimates are absent.",
    "notice_identifier": "Source notice UUID, distinct from publication number and procurement procedure ID.",
    "notice_version": "Source version number, if present. Latest versions requested from TED.",
    "notice_type": "cn-standard or cn-social.",
    "title_language": "Language key of selected title; original content retained.",
    "buyer_language": "Language key of selected buyer names.",
    "lot_count": "Number of distinct lot identifiers present in this response; zero means unknown.",
}


def csv_value(value):
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


def write_csv(path, columns, rows):
    with Path(path).open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="raise", lineterminator="\n")
        writer.writeheader()
        writer.writerows({k: csv_value(r.get(k)) for k in columns} for r in rows)


def read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_sheet(ws, columns, rows, widths=None):
    ws.append(columns)
    widths = widths or {}
    for row in rows:
        ws.append([row.get(k) for k in columns])
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for i, name in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, max(18, min(32, len(name) + 3)))
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 32
    for row in ws.iter_rows(min_row=2):
        line_count = 1
        for cell, key in zip(row, columns):
            value = cell.value
            cell.font = Font(name="Arial", size=10, color="172B4D")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(value, str):
                # Set explicit string type after assignment, even for =,+,-,@ text.
                cell.data_type = "s"
                width = ws.column_dimensions[cell.column_letter].width
                line_count = max(line_count, math.ceil(len(value) / max(1, width - 3)))
            if key in ("publication_date", "deadline") and value:
                cell.value = date.fromisoformat(value)
                cell.number_format = "yyyy-mm-dd"
            elif key == "estimated_value" and value is not None:
                cell.value = Decimal(str(value))
                cell.number_format = "#,##0.00########"
            elif key in ("notice_id", "cpv_code", "notice_identifier"):
                cell.number_format = "@"
            if key == "source_url" and isinstance(value, str) and value.startswith("https://ted.europa.eu/"):
                cell.hyperlink = value
                cell.font = Font(name="Arial", size=10, color="155EAD", underline="single")
        ws.row_dimensions[row[0].row].height = min(409, max(30, line_count * 15))


def export_workbook(path, selected, rejected, metrics, manual, automated):
    path = Path(path)
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb.create_sheet("Notices"), COLUMNS, selected,
                {"title": 65, "buyer": 48, "cpv_description": 48, "source_url": 60, "qa_note": 55,
                 "notice_identifier": 39, "estimated_value": 25})
    summary = [{"Metric": k, "Value": v} for k, v in metrics.items()]
    summary += [{"Metric": "manual_"+k, "Value": v} for k, v in manual.items()]
    summary += [{"Metric": "automated_"+k, "Value": v} for k, v in automated.items() if k in ("status", "tests", "failures", "errors")]
    write_sheet(wb.create_sheet("QA Summary"), ["Metric", "Value"], summary, {"Metric": 38, "Value": 90})
    dictionary = [{"Column": k, "Definition": DICTIONARY[k]} for k in COLUMNS]
    write_sheet(wb.create_sheet("Data Dictionary"), ["Column", "Definition"], dictionary,
                {"Column": 26, "Definition": 110})
    write_sheet(wb.create_sheet("Rejected"), REJECT_COLUMNS, rejected,
                {"title": 65, "source_url": 60, "raw_ref": 42, "reason_code": 40, "reason_detail": 80})
    temporary = path.with_name(path.stem + ".tmp.xlsx")
    wb.save(temporary)
    wb.close()
    verify_workbook(temporary, selected)
    temporary.replace(path)


def verify_workbook(path, selected):
    wb = load_workbook(path, data_only=False)
    try:
        ws = wb["Notices"]
        if list(next(ws.values)) != COLUMNS or ws.max_row != len(selected) + 1:
            raise ValueError("Excel shape mismatch")
        for values, expected in zip(list(ws.values)[1:], selected):
            for key, actual in zip(COLUMNS, values):
                value = expected.get(key)
                if isinstance(actual, datetime):
                    actual = actual.date().isoformat()
                if key == "estimated_value" and value is not None:
                    equal = actual is not None and Decimal(str(actual)) == Decimal(str(value))
                else:
                    equal = csv_value(actual) == csv_value(value)
                if not equal:
                    raise ValueError(f"Excel value mismatch: {expected['notice_id']} {key}")
        for sheet in wb:
            for row in sheet:
                if any(cell.data_type == "f" for cell in row):
                    raise ValueError("Unexpected executable spreadsheet formula")
    finally:
        wb.close()
